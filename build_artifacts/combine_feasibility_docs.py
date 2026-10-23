from copy import deepcopy
from pathlib import Path
import re

from docx import Document

from create_feasibility_docs import base_document, cover, add_control, add_contents, table


ROOT = Path('/Users/eduardolegh/Documents/galileo/formulariosInteligentes')
OUT = ROOT / 'outputs'
PRELIMINARY = OUT / 'Formularios_Inteligentes_Investigacion_Preliminar.docx'
ECONOMIC = OUT / 'Formularios_Inteligentes_Factibilidad_Economica.docx'
COMBINED = OUT / 'Formularios_Inteligentes_Estudio_de_Factibilidad.docx'


def paragraph_text(element):
    return ''.join(node.text or '' for node in element.iter() if node.tag.endswith('}t'))


def paragraph_style(element):
    ppr = element.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
    if ppr is None:
        return ''
    style = ppr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pStyle')
    if style is None:
        return ''
    return style.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '')


def replace_paragraph_text(element, text):
    text_nodes = [node for node in element.iter() if node.tag.endswith('}t')]
    if not text_nodes:
        return
    text_nodes[0].text = text
    for node in text_nodes[1:]:
        node.text = ''


def normalize_terms(doc):
    replacements = [
        ('la primera versión de Formularios Inteligentes', 'el sistema Formularios Inteligentes'),
        ('aplicación web low-code', 'sistema web'),
        ('MVP web con React', 'Sistema web con React'),
        ('OpenAI API', 'Interfaz de programación de aplicaciones de OpenAI'),
        ('JSON Schema', 'un esquema estructurado de datos'),
        ('Row Level Security', 'seguridad a nivel de filas'),
        ('Edge Functions', 'funciones seguras del servidor'),
        ('SaaS', 'software como servicio'),
        ('REST genérica', 'servicios web genéricos'),
        ('SLA empresarial', 'acuerdo empresarial de nivel de servicio'),
        ('SSO', 'inicio de sesión único'),
        ('MAU', 'usuarios activos mensuales'),
        ('GB', 'gigabytes'),
    ]
    for node in doc.element.iter():
        if not node.tag.endswith('}t') or not node.text:
            continue
        text = node.text
        for old, new in replacements:
            text = text.replace(old, new)
        text = re.sub(r'\bMVP\b', 'sistema', text)
        text = re.sub(r'\bIA\b', 'inteligencia artificial', text)
        text = re.sub(r'\bRLS\b', 'seguridad a nivel de filas', text)
        text = re.sub(r'\bSQL\b', 'base de datos', text)
        text = re.sub(r'USD\s*([0-9][0-9,]*(?:\.[0-9]+)?)', r'\1 dólares estadounidenses', text)
        text = text.replace('Tarifa USD', 'Tarifa en dólares')
        text = text.replace('Subtotal USD', 'Subtotal en dólares')
        text = text.replace('Costo USD', 'Costo en dólares')
        text = text.replace('Costo teórico USD', 'Costo teórico en dólares')
        text = text.replace('Reserva recomendada USD', 'Reserva recomendada en dólares')
        text = text.replace('Beneficio mensual USD', 'Beneficio mensual en dólares')
        text = text.replace('Desarrollo USD', 'Desarrollo en dólares')
        text = text.replace('Operación año 1 USD', 'Operación del año 1 en dólares')
        text = text.replace('Contingencia USD', 'Contingencia en dólares')
        text = text.replace('Total año 1 USD', 'Total del año 1 en dólares')
        text = text.replace('Infraestructura USD', 'Infraestructura en dólares')
        text = text.replace('Costo USD', 'Costo en dólares')
        text = text.replace('Referencia mensual USD', 'Referencia mensual en dólares')
        text = text.replace('Prototipo USD mes', 'Prototipo en dólares por mes')
        text = text.replace('Producción base USD mes', 'Producción base en dólares por mes')
        text = text.replace('Beneficio neto mes USD', 'Beneficio neto mensual en dólares')
        text = text.replace('faltan SLA,', 'faltan acuerdos de nivel de servicio,')
        text = text.replace('una sistema web', 'un sistema web')
        text = text.replace('Análisis inteligencia artificial', 'Análisis de inteligencia artificial')
        text = text.replace('análisis inteligencia artificial', 'análisis de inteligencia artificial')
        text = text.replace('Funciones inteligencia artificial', 'Funciones de inteligencia artificial')
        text = text.replace('funciones inteligencia artificial', 'funciones de inteligencia artificial')
        text = text.replace('políticas seguridad', 'políticas de seguridad')
        text = text.replace('migraciones base de datos', 'migraciones de base de datos')
        text = text.replace('integración servicios web', 'integración con servicios web')
        text = text.replace('suscripción software como servicio', 'suscripción de software como servicio')
        text = text.replace(
            'sistema en evolución; faltan acuerdos de nivel de servicio, integraciones amplias y madurez operativa.',
            'Requiere acuerdos de nivel de servicio, integraciones amplias y mayor madurez operativa.',
        )
        node.text = text


def copy_content(source_path, target_doc, start_heading, number_offset=0):
    source = Document(source_path)
    copying = False
    for element in source.element.body.iterchildren():
        if element.tag.endswith('}sectPr'):
            continue
        text = paragraph_text(element).strip() if element.tag.endswith('}p') else ''
        if not copying and text == start_heading:
            copying = True
        if not copying:
            continue

        clone = deepcopy(element)
        if number_offset and clone.tag.endswith('}p') and paragraph_style(clone).startswith('Heading'):
            heading = paragraph_text(clone).strip()
            match = re.match(r'^(\d+)(\.\d+)?\s+(.*)$', heading)
            if match:
                major = int(match.group(1)) + number_offset
                minor = match.group(2) or ''
                replace_paragraph_text(clone, f'{major}{minor} {match.group(3)}')
        target_doc.element.body.insert(-1, clone)


def copy_preliminary_as_narrative(source_path, target_doc):
    transitions = {
        '3.1 Causas principales': 'Las principales causas identificadas son las siguientes:',
        '3.2 Efectos observables': 'Estas condiciones producen efectos que pueden observarse en la operación diaria:',
        '3.3 Pregunta orientadora': 'A partir de esta situación, la investigación se orienta mediante la siguiente pregunta:',
        '5 Usuarios y necesidades': 'La solución debe responder a las necesidades de distintos participantes del proceso, desde quienes administran la plataforma hasta quienes completan y revisan los formularios.',
        '6.1 Capacidades incluidas en el MVP': 'Estas son las funcionalidades disponibles en el sistema:',
        '7 Investigación de soluciones existentes': 'La revisión del mercado confirma que ya existen herramientas capaces de resolver una parte del problema. Las alternativas más relevantes se describen a continuación.',
        '8 Análisis competitivo': 'La comparación permite observar las fortalezas de cada solución y las limitaciones que presentan frente al flujo propuesto por Formularios Inteligentes.',
        '9 Oportunidad y riesgos preliminares': 'Además de la oportunidad funcional, el proyecto debe considerar los riesgos de seguridad, calidad de la IA, dependencia de proveedores y adopción por parte de los usuarios.',
    }
    source = Document(source_path)
    copying = False
    for element in source.element.body.iterchildren():
        if element.tag.endswith('}sectPr'):
            continue
        text = paragraph_text(element).strip() if element.tag.endswith('}p') else ''
        if not copying and text == '2 Situación actual y contexto':
            copying = True
        if not copying:
            continue

        if text.startswith('El caso demostrativo del proyecto es el reclutamiento de personal.'):
            continue
        if text.startswith('El MVP no pretende reemplazar plataformas empresariales completas'):
            continue

        if element.tag.endswith('}p') and paragraph_style(element).startswith('Heading'):
            if text == '11 Fuentes consultadas':
                p = target_doc.add_paragraph()
                p.paragraph_format.space_before = None
                p.add_run('Fuentes consultadas').bold = True
            elif text in transitions:
                target_doc.add_paragraph(transitions[text])
            continue

        if element.tag.endswith('}tbl'):
            table_text = paragraph_text(element).strip()
            if table_text.startswith('MotivaciónAplicación en el proyecto'):
                target_doc.add_paragraph(
                    'El proyecto busca reducir el esfuerzo necesario para configurar formularios mediante '
                    'su generación a partir de instrucciones escritas en lenguaje natural y estructuras de '
                    'datos validadas. También pretende mejorar la consistencia al permitir que cada formulario '
                    'incluya un contexto definido por objetivos, instrucciones, reglas, restricciones y '
                    'resultados esperados. La protección de la información se garantiza mediante mecanismos '
                    'de autenticación, roles, permisos y políticas de seguridad en Supabase. Además, el sistema '
                    'agiliza la revisión de las respuestas utilizando inteligencia artificial para proponer un '
                    'estado y una observación, aunque la decisión debe ser confirmada por una persona antes de '
                    'guardarse. La trazabilidad se mantiene mediante los estados Ingresado, En revisión, '
                    'Aprobado y Rechazado, junto con el registro permanente de las observaciones. Finalmente, '
                    'la arquitectura web modular, las migraciones de la base de datos y las funciones seguras '
                    'del servidor facilitan el mantenimiento y la incorporación de mejoras en el sistema.'
                )
                continue
            if table_text.startswith('UsuarioNecesidad principalRespuesta de la solución'):
                target_doc.add_paragraph(
                    'El sistema atiende las necesidades de diferentes tipos de usuarios. El administrador '
                    'controla los usuarios, roles, departamentos y la operación general mediante un directorio '
                    'de usuarios y herramientas para asignar funciones y permisos administrativos. El diseñador '
                    'de formularios puede crear y modificar formularios mediante un editor visual y funciones '
                    'de generación asistida por inteligencia artificial, lo que reduce el trabajo repetitivo. '
                    'Por su parte, el configurador de contexto establece los criterios de negocio que orientan '
                    'el análisis, incluyendo reglas, restricciones e instrucciones que pueden actualizarse '
                    'conforme cambien las necesidades del proceso.'
                )
                target_doc.add_paragraph(
                    'Los participantes pueden responder los formularios desde diferentes dispositivos mediante '
                    'enlaces públicos o privados, con validación de campos y captura de imágenes cuando sea '
                    'necesario. Los usuarios encargados de la revisión consultan únicamente los envíos '
                    'correspondientes a su departamento mediante una bandeja filtrada, desde la cual pueden '
                    'examinar las respuestas, solicitar el análisis mediante inteligencia artificial, registrar '
                    'observaciones y tomar una decisión. Finalmente, el analista puede consultar el volumen de '
                    'solicitudes, sus estados y resultados a partir de información preparada para la creación '
                    'de paneles e indicadores.'
                )
                continue

        clone = deepcopy(element)
        # Las listas conservan su contenido, pero se presentan como parte del
        # discurso continuo y no como subsecciones independientes.
        target_doc.element.body.insert(-1, clone)


def copy_economic_as_narrative(source_path, target_doc):
    transitions = {
        '2.1 Criterios utilizados': 'La estimación se elaboró con los siguientes criterios:',
        '3 Supuestos de estimación': 'Para calcular los costos se definieron los siguientes supuestos de uso y operación:',
        '4 Recursos necesarios': 'La implementación y operación del sistema requieren recursos humanos y tecnológicos.',
        '4.1 Recursos humanos': 'El principal recurso humano corresponde al tiempo especializado del equipo de trabajo.',
        '4.2 Recursos tecnológicos': 'Los recursos tecnológicos considerados se describen a continuación:',
        '5 Costos de desarrollo': 'El costo de desarrollo se determinó a partir de las horas estimadas para cada actividad.',
        '5.1 Distribución porcentual': 'La inversión de desarrollo se distribuye porcentualmente de la siguiente manera:',
        '5.2 Costos no incluidos': 'La estimación no incorpora determinados gastos que dependerán del uso institucional del sistema.',
        '6 Costos de infraestructura y operación': 'Los costos de infraestructura y operación dependen del entorno utilizado y del consumo mensual.',
        '7 Costo de inteligencia artificial': 'El uso de inteligencia artificial genera un costo variable asociado con la cantidad de información procesada.',
        '7.1 Estimación por operación': 'El costo aproximado de cada análisis se calculó con los siguientes supuestos:',
        '8 Escenarios presupuestarios': 'A partir de los costos identificados se plantean distintos escenarios presupuestarios:',
        '8.1 Flujo mensual después del lanzamiento': 'Después del lanzamiento, el flujo mensual combina infraestructura y mantenimiento.',
        '9 Beneficios y retorno estimado': 'La factibilidad también considera los beneficios operativos y el tiempo estimado para recuperar la inversión.',
        '9.1 Escenario de beneficio': 'El escenario esperado de beneficio se basa en el ahorro de tiempo durante la revisión, creación y seguimiento.',
        '9.2 Sensibilidad': 'El período de recuperación cambia según la cantidad mensual de envíos, como muestra el análisis de sensibilidad.',
        '10 Comparación económica con alternativas': 'La comparación con soluciones existentes permite valorar el costo frente al nivel de control y personalización ofrecido.',
        '11 Riesgos económicos y medidas de control': 'La estimación contempla riesgos económicos que deben administrarse mediante medidas de control.',
    }
    source = Document(source_path)
    copying = False
    for element in source.element.body.iterchildren():
        if element.tag.endswith('}sectPr'):
            continue
        text = paragraph_text(element).strip() if element.tag.endswith('}p') else ''
        if not copying and text == '2 Objetivo y metodología':
            copying = True
        if not copying:
            continue

        if element.tag.endswith('}p') and paragraph_style(element).startswith('Heading'):
            if text == '2 Objetivo y metodología' or text == '2.2 Alcance de la estimación':
                continue
            if text == 'Nota sobre la estimación':
                continue
            if text in ('12 Conclusión de factibilidad', '13 Fuentes consultadas'):
                labels = {
                    '12 Conclusión de factibilidad': 'Conclusión de factibilidad',
                    '13 Fuentes consultadas': 'Fuentes consultadas',
                }
                p = target_doc.add_paragraph()
                p.add_run(labels[text]).bold = True
            elif text in transitions:
                target_doc.add_paragraph(transitions[text])
            continue

        target_doc.element.body.insert(-1, deepcopy(element))


def add_economic_feasibility(target_doc):
    target_doc.add_paragraph(
        'Esta factibilidad económica constituye una primera estimación de los costos del proyecto. En ella se '
        'desglosan los recursos humanos, tecnológicos y operativos necesarios. Los valores son referenciales y '
        'deberán refinarse después del análisis y diseño. Las horas se valoran a 15 dólares estadounidenses por '
        'programador y no representan un salario, una cotización comercial ni un pago inmediato.'
    )
    target_doc.add_paragraph(
        'El trabajo se realizará entre dos programadores con responsabilidades compartidas. Para esta primera '
        'estimación se considera una carga total de 380 horas: 190 horas para cada integrante. El primer '
        'programador tendrá mayor participación en la interfaz, mientras que el segundo se concentrará en el '
        'servidor y la base de datos. Las tareas de análisis, diseño, inteligencia artificial, pruebas, despliegue '
        'y documentación se distribuirán de forma equilibrada.'
    )
    table(target_doc, ['Actividad', 'Programador 1', 'Programador 2', 'Total', 'Costo estimado'], [
        ['Análisis y definición del alcance', '20 h', '20 h', '40 h', '600 dólares'],
        ['Diseño de experiencia y formularios', '20 h', '20 h', '40 h', '600 dólares'],
        ['Desarrollo de la interfaz', '60 h', '40 h', '100 h', '1,500 dólares'],
        ['Servidor y base de datos', '40 h', '60 h', '100 h', '1,500 dólares'],
        ['Inteligencia artificial y pruebas', '30 h', '30 h', '60 h', '900 dólares'],
        ['Despliegue y documentación', '20 h', '20 h', '40 h', '600 dólares'],
        ['Total estimado', '190 h', '190 h', '380 h', '5,700 dólares'],
    ], widths=[2.4, 1.05, 1.05, 0.8, 1.35], font_size=7.8)
    target_doc.add_paragraph(
        'El trabajo se valora en 5,700 dólares estadounidenses, equivalentes a 2,850 dólares por programador. '
        'Con una dedicación aproximada de 15 horas semanales por persona, el desarrollo requeriría alrededor de '
        '13 semanas. Estas cifras reconocen el esfuerzo académico, pero no constituyen un desembolso inmediato.'
    )
    target_doc.add_paragraph(
        'React, TypeScript, Vite y Git no requieren licencias. Para el escenario académico se utilizarán los '
        'niveles gratuitos de Supabase y Vercel, junto con una reserva mensual de 2 dólares para inteligencia '
        'artificial. El dominio propio es opcional y se estima en 15 dólares al año.'
    )

    table(target_doc, ['Recurso', 'Costo estimado', 'Justificación'], [
        ['Equipo de desarrollo', '0 dólares', 'Se utilizan computadoras disponibles del equipo.'],
        ['Herramientas de programación', '0 dólares', 'React, TypeScript, Vite y Git son herramientas sin costo de licencia.'],
        ['Base de datos y autenticación', '0 dólares al mes', 'El nivel gratuito de Supabase cubre el uso académico previsto.'],
        ['Alojamiento del sistema', '0 dólares al mes', 'El nivel gratuito de Vercel es suficiente para pruebas y demostraciones.'],
        ['Inteligencia artificial', '2 dólares al mes', 'Reserva para generación de formularios y análisis de respuestas.'],
        ['Dominio propio', '15 dólares al año, opcional', 'Solo se necesita si se desea una dirección personalizada.'],
    ], widths=[2.0, 1.5, 3.2])

    target_doc.add_paragraph(
        'El desembolso operativo se estima en 24 dólares anuales sin dominio o 39 dólares con dominio propio. '
        'No se requiere adquirir licencias, servidores ni equipo adicional durante el escenario académico.'
    )
    target_doc.add_paragraph(
        'Inicialmente se utilizarán modelos de OpenAI. Posteriormente se evaluará Ollama mediante pruebas de '
        'rendimiento, calidad de respuesta y consumo de recursos. Los modelos locales podrían reducir la dependencia '
        'de servicios externos, pero también pueden requerir equipos con mayor capacidad, energía y mantenimiento; '
        'por ello, no se incluyen como gasto inmediato.'
    )
    target_doc.add_paragraph(
        'Si el sistema pasa a un uso institucional, deberán revisarse los límites gratuitos. Como referencia, '
        'Supabase ofrece un plan profesional desde 25 dólares mensuales y Vercel desde 20 dólares mensuales; '
        'estos planes no forman parte del presupuesto actual.'
    )
    table(target_doc, ['Riesgo económico', 'Medida de control'], [
        ['Aumento del consumo de inteligencia artificial', 'Impedir análisis repetidos, registrar el consumo y establecer un límite mensual.'],
        ['Superación de los límites gratuitos', 'Revisar trimestralmente usuarios, almacenamiento y transferencias.'],
        ['Cambio de precios de proveedores', 'Actualizar el presupuesto antes de contratar cualquier plan pagado.'],
        ['Rendimiento insuficiente de modelos locales', 'Comparar Ollama con los modelos de OpenAI antes de realizar la migración.'],
        ['Crecimiento no planificado del sistema', 'Autorizar nuevas funciones según prioridad y disponibilidad presupuestaria.'],
    ], widths=[2.7, 4.0])

    p = target_doc.add_paragraph()
    p.add_run('Conclusión de factibilidad. ').bold = True
    p.add_run(
        'El proyecto es económicamente factible. La primera estimación reconoce 5,700 dólares estadounidenses '
        'como valor del trabajo del equipo y entre 24 y 39 dólares como desembolso operativo durante el primer '
        'año, según se adquiera o no un dominio propio. Por tanto, el costo económico total de referencia se '
        'encuentra entre 5,724 y 5,739 dólares, aunque el pago efectivo previsto para el escenario académico '
        'corresponde únicamente a la operación. El uso de infraestructura gratuita reduce el desembolso sin '
        'impedir la demostración de las funciones principales. Todas estas cifras deberán refinarse después de '
        'la fase de análisis y diseño y conforme se obtengan datos reales de uso.'
    )
    p = target_doc.add_paragraph()
    p.add_run('Fuentes consultadas. ').bold = True
    p.add_run(
        'Precios y características oficiales de Supabase, Vercel y OpenAI, consultados el 3 de octubre de 2026. '
        'Los valores son referencias públicas y deben verificarse nuevamente antes de contratar servicios.'
    )


def build():
    doc = base_document('Estudio de factibilidad')
    cover(
        doc,
        'Estudio de Factibilidad',
        'Investigación preliminar y factibilidad económica de Formularios Inteligentes',
        'FI EF 01',
    )
    add_control(doc, 'Estudio de factibilidad: investigación preliminar y factibilidad económica')
    add_contents(doc, [
        'Presentación del estudio',
        'Investigación preliminar desarrollada de forma narrativa',
        'Factibilidad económica',
        'Objetivo, metodología y supuestos',
        'Recursos y costos del proyecto',
        'Escenarios, beneficios y riesgos',
        'Conclusión y fuentes consultadas',
    ])

    doc.add_heading('Presentación del estudio', level=1)
    doc.add_paragraph(
        'Este documento integra la investigación preliminar y la factibilidad económica del proyecto '
        'Formularios Inteligentes. La primera parte analiza el contexto, el problema, las motivaciones, '
        'los usuarios, la solución propuesta y sus principales competidores. La segunda parte estima '
        'los recursos, costos, escenarios y beneficios necesarios para implementar y operar el producto.'
    )
    doc.add_paragraph(
        'En conjunto, el estudio permite valorar la pertinencia de la solución y establecer una base '
        'económica inicial para las fases de análisis, diseño e implementación. Las cifras corresponden '
        'a una estimación preliminar y deberán actualizarse con los resultados del piloto y las decisiones '
        'definitivas de arquitectura y alcance.'
    )

    part_one = doc.add_heading('Parte I Investigación preliminar', level=1)
    part_one.paragraph_format.page_break_before = True
    part_one.paragraph_format.space_before = None
    doc.add_paragraph(
        'La investigación preliminar muestra que el mercado ya ofrece herramientas maduras para crear '
        'formularios y recopilar respuestas. La oportunidad de Formularios Inteligentes consiste en reunir '
        'en una sola plataforma el contexto del negocio, los permisos por departamento, la revisión de '
        'solicitudes, el análisis asistido por inteligencia artificial y el registro de decisiones. El '
        'proyecto resulta especialmente pertinente para organizaciones que actualmente combinan formularios '
        'genéricos, hojas de cálculo, correo electrónico y tareas manuales de seguimiento.'
    )
    copy_preliminary_as_narrative(PRELIMINARY, doc)

    part_two = doc.add_heading('Parte II Factibilidad económica', level=1)
    part_two.paragraph_format.page_break_before = True
    part_two.paragraph_format.space_before = None
    doc.add_paragraph(
        'La aplicación es económicamente viable porque puede utilizar infraestructura gratuita durante su '
        'etapa académica y requiere únicamente una reserva pequeña para el consumo de inteligencia artificial. '
        'La estimación distingue el valor referencial de las horas de trabajo y los desembolsos que efectivamente '
        'deberá realizar el equipo. Las cifras se actualizarán después de la fase de análisis y diseño.'
    )
    add_economic_feasibility(doc)

    doc.core_properties.title = 'Estudio de Factibilidad de Formularios Inteligentes'
    doc.core_properties.subject = 'Investigación preliminar y factibilidad económica'
    doc.core_properties.author = 'Equipo de Formularios Inteligentes'
    normalize_terms(doc)
    doc.save(COMBINED)
    return COMBINED


if __name__ == '__main__':
    print(build())
